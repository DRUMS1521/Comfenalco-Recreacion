"""
Migración del Cronograma de Recreadores (Excel, 52 hojas semanales) a Solicitud.

Estructura del Excel: una hoja por semana, con nombre "DEL X AL Y DE MES".
Fila 2: encabezado de día ("5 LUNES", "6 MARTES", ...) en la primera columna
de cada bloque de 3 (EVENTO, HORA LABORAL, HORA FLEXIBLE). Filas 4 en
adelante: una fila por recreador, con el nombre en la columna A.

ADVERTENCIA: esta migración es "todo sin filtrar" (decisión explícita del
usuario) — cada celda EVENTO puede contener varios bloques de texto
separados por "----", y TODOS se migran como Solicitud, sin distinguir
entre eventos de recreación reales (pasadías, planes promocionales) y
tareas administrativas internas (exámenes médicos, capacitaciones,
bodega, incapacidades). El resultado es un calendario con mucho más
volumen y ruido que si se hubiera filtrado.

Antes de insertar, borra TODAS las solicitudes existentes (no toca
usuarios ni empresas).
"""
import sys, os, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from datetime import date, datetime, timezone
from collections import defaultdict
import openpyxl

from app.core.database import SessionLocal, engine, Base
from app.models.solicitud import Solicitud, solicitud_recreadores
from app.models.user import User
from sqlalchemy import text

Base.metadata.create_all(bind=engine)

EXCEL_PATH = "/Users/drums/Downloads/CRONOGRAMA RECREADORES 08 DE SEPTIEMBRE 2026.xlsx"
HOY = date(2026, 9, 9)
CREADOR_USERNAME = "jennifer"

MESES = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
    "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11,
    "DICIEMBRE": 12,
}
MESES_NOMBRES = list(MESES.keys())

NOMBRE_A_USERNAME = {
    "GABRIEL VAQUIRO": "gabriel.vaquiro",
    "ANDRES GARCIA": "andres.garcia",
    "SEBASTIAN LEGUIZAMON": "sebastian.leguizamon",
    "DALMA SUAREZ": "dalma.suarez",
    "JOHNY FANDIÑO -GUIA": "johny.fandino",
    "JOHNY FANDIÑO": "johny.fandino",
    "MERCEDES LEAL": "mercedes.leal",
    "NESTOR VILLA": "nestor.villa",
    "ANGIE ROJAS": "angie.rojas",
    "CAMILO PEÑA": "camilo.pena",
    "ISSA PAEZ": "issa.paez",
    "JUAN DANIEL": "juan.daniel",
    "MAURICIO ARIAS": "mauricio.arias",
    "DANIEL RINCON": "daniel.rincon",
    "DANIEL RUIZ": "daniel.ruiz",
    "BRAYAN BOTERO": "brayan.botero",
    "BAIRON ROJAS": "bairon.rojas",
}

DAY_COLUMNS = [4, 7, 10, 13, 16, 19, 22]  # D, G, J, M, P, S, V (1-based)


# ---------------------------------------------------------------- fechas ---

def extraer_meses(semana: str):
    tokens = re.findall(r"[A-ZÑ]+", semana.upper())
    encontrados = []
    for t in tokens:
        if len(t) < 3:
            continue
        for mes in MESES_NOMBRES:
            if mes.startswith(t) or t.startswith(mes[:4]):
                encontrados.append(mes)
                break
    return encontrados


def extraer_dias_rango(semana: str):
    nums = [int(n) for n in re.findall(r"\d+", semana)]
    if not nums:
        return None, None
    return nums[0], nums[-1]


def extraer_dia_numero(dia: str):
    m = re.search(r"\d+", dia or "")
    return int(m.group()) if m else None


def resolver_fecha(semana: str, dia: str):
    meses = extraer_meses(semana)
    start_day, _ = extraer_dias_rango(semana)
    d = extraer_dia_numero(dia)
    if not meses or d is None:
        return None
    mes_inicio, mes_fin = meses[0], meses[-1]
    mes = mes_inicio if mes_inicio == mes_fin else (
        mes_inicio if (start_day is not None and d >= start_day) else mes_fin
    )
    try:
        return date(2026, MESES[mes], d).isoformat()
    except ValueError:
        return None


# ------------------------------------------------------------- horarios ---

HORA_RE = re.compile(
    r"(\d{1,2}:\d{2})\s*([ap]\.?\s*m\.?|M)\b.{0,10}?\b[aA]\b\s*(\d{1,2}:\d{2})\s*([ap]\.?\s*m\.?|M)\b",
)


def _a_24h(hhmm, ampm):
    h, m = map(int, hhmm.split(":"))
    ampm = ampm.strip().lower().replace(".", "")
    if ampm == "m":
        pass  # "12:00 M" == mediodía == 12:00 PM, ya en 24h
    elif ampm == "pm" and h != 12:
        h += 12
    elif ampm == "am" and h == 12:
        h = 0
    return f"{h:02d}:{m:02d}"


def extraer_horario(texto: str):
    matches = HORA_RE.findall(texto)
    if not matches:
        return "08:00", "17:00"
    ini = _a_24h(matches[0][0], matches[0][1])
    fin = _a_24h(matches[-1][2], matches[-1][3])
    return ini, fin


def horas_bloques_del_dia(n_bloques: int, horas_reales: float):
    """
    Reparte las horas REALES del día (columna HORA LABORAL + HORA FLEXIBLE
    del Excel, la fuente de verdad) en partes iguales entre los N bloques
    de texto de ese día, como intervalos contiguos desde las 08:00.

    Esto evita inflar el total semanal: antes se asignaba un horario fijo
    08:00–17:00 (9h) a CADA bloque, así que un día con 3 bloques sumaba
    27h aunque el Excel dijera que ese día se trabajaron, por ejemplo,
    solo 6h.
    """
    if n_bloques <= 0:
        return []
    share = (horas_reales or 0) / n_bloques
    horarios = []
    cursor = 0.0
    for _ in range(n_bloques):
        ini_min = round(cursor * 60)
        fin_min = round((cursor + share) * 60)
        ini_h, ini_m = divmod(ini_min, 60)
        fin_h, fin_m = divmod(fin_min, 60)
        horarios.append((f"{(8 + ini_h) % 24:02d}:{ini_m:02d}", f"{(8 + fin_h) % 24:02d}:{fin_m:02d}"))
        cursor += share
    return horarios


# --------------------------------------------------------- otros campos ---

def extraer_personas(texto: str):
    m = re.search(r"(?:No\.?\s*[Dd]e\s*personas|PAX|cantidad)\D{0,5}(\d+)", texto, re.IGNORECASE)
    return int(m.group(1)) if m else None


def extraer_empresa(texto: str):
    m = re.search(r"EMPRESA\s+([A-ZÑÁÉÍÓÚ0-9 ]+?)(?:,|\.|$|NIT)", texto.upper())
    if m:
        return m.group(1).strip().title()
    return "Comfenalco Tolima - Cronograma Interno"


def extraer_contacto(texto: str):
    contacto, telefono = None, None
    m = re.search(r"(?:USUARIO|CONTACTO(?: EMPRESA)?)[.:]?\s*([A-Za-zÑñÁÉÍÓÚáéíóú .]{4,40}?)(?:,|CEL|$)", texto, re.IGNORECASE)
    if m:
        contacto = m.group(1).strip()
    m2 = re.search(r"(?:CEL(?:ULAR)?)[.:]?\s*([\d ]{7,})", texto, re.IGNORECASE)
    if m2:
        telefono = m2.group(1).strip()
    return contacto or "Comfenalco Tolima (interno)", telefono or "N/A"


def clasificar_tipo(texto: str):
    t = texto.upper()
    if "PLAN PROMOCIONAL" in t:
        return "Plan promocional"
    if "PASADIA" in t or "PASADÍA" in t:
        return "Pasadía"
    if "INCAPACIDAD" in t:
        return "Incapacidad / novedad"
    if "EXAMEN" in t or "EXÁMEN" in t or "IPS" in t:
        return "Trámite administrativo"
    if "CAPACITACION" in t or "CAPACITACIÓN" in t or "TALLER" in t:
        return "Capacitación interna"
    if "BODEGA" in t or "ARREGLO" in t or "DECORACION" in t:
        return "Tarea administrativa"
    if "ESCUELA" in t or "ESCUELAS DEPORTIVAS" in t:
        return "Apoyo escuela deportiva"
    if "FIESTA" in t:
        return "Evento social"
    if "ACTIVIDAD RECREATIVA" in t or "CARDIO RUMBA" in t:
        return "Actividad recreativa"
    if "VIAJE" in t:
        return "Viaje / traslado"
    return "Actividad interna"


def ciudad_de(texto: str):
    m = re.search(r"\b(?:A|DESDE)\s+([A-ZÑÁÉÍÓÚ][A-ZÑÁÉÍÓÚ .]{2,30})", texto.upper())
    if m:
        destino = m.group(1).strip().rstrip(".,")
        destino = re.split(r"\s{2,}|,|\.", destino)[0].strip()
        return destino.title()
    return "Ibagué"


# -------------------------------------------------------------- migrar ---

def migrar():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    db = SessionLocal()
    try:
        creador = db.query(User).filter(User.username == CREADOR_USERNAME).first()
        if not creador:
            print(f"✗ No se encontró el usuario creador '{CREADOR_USERNAME}'.")
            return

        # 1. Limpiar solicitudes existentes (usuarios y empresas quedan intactos)
        borradas = db.query(Solicitud).count()
        db.execute(text("DELETE FROM solicitud_recreadores"))
        db.query(Solicitud).delete()
        db.commit()
        print(f"🗑  {borradas} solicitudes previas eliminadas.")

        creados = 0
        nombres_no_encontrados = set()
        dias_no_resueltos = defaultdict(int)
        por_hoja = {}

        recreadores_cache = {u.username: u for u in db.query(User).filter(User.is_recreador == True).all()}  # noqa: E712

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            creados_hoja = 0

            # Resolver fecha de cada columna de día
            fechas_col = {}
            for col in DAY_COLUMNS:
                label = ws.cell(row=2, column=col).value
                if not label:
                    continue
                fecha = resolver_fecha(sheet_name, str(label))
                if fecha:
                    fechas_col[col] = fecha
                else:
                    dias_no_resueltos[f"{sheet_name} / {label}"] += 1

            for row in range(4, ws.max_row + 1):
                nombre_raw = ws.cell(row=row, column=1).value
                if not nombre_raw or not str(nombre_raw).strip():
                    continue
                nombre = str(nombre_raw).strip().upper()
                username = NOMBRE_A_USERNAME.get(nombre)
                if not username:
                    nombres_no_encontrados.add(nombre)
                    continue
                recreador = recreadores_cache.get(username)
                if not recreador:
                    nombres_no_encontrados.add(f"{nombre} (sin cuenta activa: {username})")
                    continue

                for col, fecha in fechas_col.items():
                    evento_val = ws.cell(row=row, column=col).value
                    if not evento_val or not str(evento_val).strip():
                        continue
                    bloques = [b.strip() for b in re.split(r"-{3,}", str(evento_val)) if b.strip()]
                    if not bloques:
                        continue

                    labor = ws.cell(row=row, column=col + 1).value
                    flex = ws.cell(row=row, column=col + 2).value
                    horas_reales = (labor if isinstance(labor, (int, float)) else 0) + \
                                   (flex if isinstance(flex, (int, float)) else 0)
                    horarios = horas_bloques_del_dia(len(bloques), horas_reales)

                    for bloque, (ini, fin) in zip(bloques, horarios):
                        personas = extraer_personas(bloque) or 20
                        contacto, telefono = extraer_contacto(bloque)
                        empresa = extraer_empresa(bloque)
                        tipo_servicio = clasificar_tipo(bloque)
                        ciudad = ciudad_de(bloque)
                        estado = "finalizado" if fecha <= HOY.isoformat() else "programado"

                        sol = Solicitud(
                            empresa=empresa,
                            fecha_evento=fecha,
                            hora_inicio=ini,
                            hora_fin=fin,
                            ciudad=ciudad,
                            direccion=f"{ciudad} — cronograma Comfenalco",
                            cantidad_recreadores=1,
                            cantidad_personas=personas,
                            tipo_publico="Mixto",
                            tipo_servicio=tipo_servicio,
                            contacto=contacto,
                            telefono_email=telefono,
                            observaciones=f"[{sheet_name.strip()} · {ws.cell(row=2, column=col).value}] {bloque}",
                            estado=estado,
                            recreador_id=recreador.id,
                            user_id=creador.id,
                        )
                        db.add(sol)
                        db.flush()
                        sol.recreadores = [recreador]
                        creados += 1
                        creados_hoja += 1

            por_hoja[sheet_name] = creados_hoja

        db.commit()

        print(f"\n✓ {creados} solicitudes creadas a partir del cronograma Excel (52 hojas).")
        if nombres_no_encontrados:
            print(f"\n⚠ Nombres sin cuenta de recreador coincidente ({len(nombres_no_encontrados)}):")
            for n in sorted(nombres_no_encontrados):
                print(f"   - {n}")
        if dias_no_resueltos:
            print(f"\n⚠ Columnas de día sin fecha resoluble ({len(dias_no_resueltos)}):")
            for d, c in sorted(dias_no_resueltos.items()):
                print(f"   - {d} ({c}x)")

        print("\n--- Resumen por hoja (top 5 con más registros) ---")
        for nombre, n in sorted(por_hoja.items(), key=lambda x: -x[1])[:5]:
            print(f"   {nombre.strip():<35} {n}")

    finally:
        db.close()


if __name__ == "__main__":
    migrar()
